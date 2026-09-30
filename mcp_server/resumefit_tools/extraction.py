"""Detect the type of an uploaded resume and extract its text within fixed caps.

Every function here is pure and synchronous; the worker module runs it in a
separate, killable process so a hostile file cannot stall the service.
"""

from __future__ import annotations

import io
import zipfile
from pathlib import PurePath
from typing import Any, Literal

from docx import Document
from pypdf import PdfReader

FileType = Literal["pdf", "docx", "txt"]

MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_PDF_PAGES = 5
MAX_DOCX_ENTRIES = 500
MAX_DOCX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
MAX_TEXT_CHARS = 50_000

PDF_SIGNATURE = b"%PDF-"
ZIP_SIGNATURE = b"PK\x03\x04"
WORD_DOCUMENT_PART = "word/document.xml"
ZIP_ENCRYPTED_FLAG = 0x01
NESTED_ARCHIVE_SUFFIXES = (".zip", ".docx", ".docm", ".xlsx", ".xlsm", ".pptx", ".jar")
EXTENSION_TYPES: dict[str, FileType] = {".pdf": "pdf", ".docx": "docx", ".txt": "txt"}


class ExtractionError(Exception):
    """An expected rejection, carrying one of the documented error codes."""

    def __init__(self, code: str) -> None:
        super().__init__(code)
        self.code = code


def error_result(code: str) -> dict[str, Any]:
    return {"ok": False, "error_code": code}


def extract_text(filename: str, data: bytes) -> dict[str, Any]:
    """Return the documented success or error dictionary for one uploaded file."""
    try:
        file_type = _checked_type(filename, data)
        pages: int | None = None
        if file_type == "pdf":
            text, pages = _pdf_text(data)
        elif file_type == "docx":
            text = _docx_text(data)
        else:
            text = _plain_text(data)
    except ExtractionError as error:
        return error_result(error.code)
    return {
        "ok": True,
        "file_type": file_type,
        "text": text[:MAX_TEXT_CHARS],
        "pages": pages,
        "truncated": len(text) > MAX_TEXT_CHARS,
    }


def _checked_type(filename: str, data: bytes) -> FileType:
    if len(data) > MAX_FILE_BYTES:
        raise ExtractionError("too_large")
    claimed = EXTENSION_TYPES.get(PurePath(filename).suffix.lower())
    if data.startswith(PDF_SIGNATURE):
        detected: FileType | None = "pdf"
    elif data.startswith(ZIP_SIGNATURE):
        detected = "docx"
    else:
        detected = "txt"
    if claimed is None or detected != claimed:
        raise ExtractionError("unsupported_type")
    return detected


def _pdf_text(data: bytes) -> tuple[str, int]:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ExtractionError("unreadable")
        pages = len(reader.pages)
        if pages > MAX_PDF_PAGES:
            raise ExtractionError("too_many_pages")
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    except ExtractionError:
        raise
    except Exception as error:
        raise ExtractionError("unreadable") from error
    return text.strip(), pages


def _check_docx_archive(data: bytes) -> None:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as error:
        raise ExtractionError("unreadable") from error
    with archive:
        entries = archive.infolist()
        if len(entries) > MAX_DOCX_ENTRIES:
            raise ExtractionError("archive_too_large")
        if sum(entry.file_size for entry in entries) > MAX_DOCX_UNCOMPRESSED_BYTES:
            raise ExtractionError("archive_too_large")
        for entry in entries:
            if entry.flag_bits & ZIP_ENCRYPTED_FLAG:
                raise ExtractionError("archive_unsafe")
            if entry.filename.lower().endswith(NESTED_ARCHIVE_SUFFIXES):
                raise ExtractionError("archive_unsafe")
        for entry in entries:
            try:
                with archive.open(entry) as member:
                    head = member.read(len(ZIP_SIGNATURE))
            except Exception as error:
                raise ExtractionError("unreadable") from error
            if head == ZIP_SIGNATURE:
                raise ExtractionError("archive_unsafe")
        if WORD_DOCUMENT_PART not in archive.namelist():
            raise ExtractionError("unsupported_type")


def _docx_text(data: bytes) -> str:
    _check_docx_archive(data)
    try:
        document = Document(io.BytesIO(data))
        lines = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                lines.extend(cell.text for cell in row.cells)
    except Exception as error:
        raise ExtractionError("unreadable") from error
    return "\n".join(line for line in lines if line.strip()).strip()


def _plain_text(data: bytes) -> str:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise ExtractionError("bad_encoding") from error
    if "\x00" in text:
        raise ExtractionError("bad_encoding")
    return text
